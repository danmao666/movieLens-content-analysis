from datetime import date

import pandas as pd
from openpyxl import load_workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from common import OUTPUT_DIR
from week2_common import connect_week2, require_week2_tables


def format_sheet(worksheet):
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.sheet_view.showGridLines = False
    for cell in worksheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center")
    for column_cells in worksheet.columns:
        letter = get_column_letter(column_cells[0].column)
        width = max(
            len(str(cell.value)) if cell.value is not None else 0
            for cell in column_cells
        )
        worksheet.column_dimensions[letter].width = min(max(width + 2, 12), 42)


def add_trend_chart(dashboard, trend_sheet):
    last_row = trend_sheet.max_row
    headers = {cell.value: cell.column for cell in trend_sheet[1]}
    chart = LineChart()
    chart.title = "DAU and 28-day baseline"
    chart.y_axis.title = "DAU"
    chart.x_axis.title = "Date"
    chart.height = 8
    chart.width = 16
    for name in ("dau", "dau_28d_avg"):
        column = headers.get(name)
        if column:
            data = Reference(
                trend_sheet,
                min_col=column,
                max_col=column,
                min_row=1,
                max_row=last_row,
            )
            chart.add_data(data, titles_from_data=True)
    categories = Reference(
        trend_sheet,
        min_col=headers["date_day"],
        min_row=2,
        max_row=last_row,
    )
    chart.set_categories(categories)
    dashboard.add_chart(chart, "A9")


def add_content_chart(dashboard, content_sheet):
    last_row = min(content_sheet.max_row, 21)
    headers = {cell.value: cell.column for cell in content_sheet[1]}
    chart = BarChart()
    chart.type = "bar"
    chart.title = "Top content by active users"
    chart.x_axis.title = "Active users"
    chart.y_axis.title = "Content"
    chart.height = 8
    chart.width = 16
    data = Reference(
        content_sheet,
        min_col=headers["active_users"],
        max_col=headers["active_users"],
        min_row=1,
        max_row=last_row,
    )
    categories = Reference(
        content_sheet,
        min_col=headers["title"],
        min_row=2,
        max_row=last_row,
    )
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(categories)
    dashboard.add_chart(chart, "J9")


def add_segment_chart(dashboard, segment_sheet):
    headers = {cell.value: cell.column for cell in segment_sheet[1]}
    chart = PieChart()
    chart.title = "User segments"
    chart.height = 8
    chart.width = 12
    data = Reference(
        segment_sheet,
        min_col=headers["user_count"],
        max_col=headers["user_count"],
        min_row=1,
        max_row=segment_sheet.max_row,
    )
    labels = Reference(
        segment_sheet,
        min_col=headers["user_segment"],
        min_row=2,
        max_row=segment_sheet.max_row,
    )
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(labels)
    dashboard.add_chart(chart, "A27")


def main() -> None:
    connection = connect_week2()
    try:
        require_week2_tables(
            connection,
            "week2_anomaly_monitoring",
            "week2_user_segments",
            "week2_cohort_retention",
            "week2_content_concentration",
        )

        overview = connection.sql("""
            WITH latest AS (
                SELECT *
                FROM week2_daily_metrics
                ORDER BY date_day DESC
                LIMIT 1
            ), last_28 AS (
                SELECT *
                FROM week2_daily_metrics
                ORDER BY date_day DESC
                LIMIT 28
            )
            SELECT
                latest.date_day AS latest_date,
                latest.dau AS latest_dau,
                latest.behavior_count AS latest_behavior_count,
                latest.active_content_count AS latest_active_content,
                ROUND((SELECT AVG(dau) FROM last_28), 2)
                    AS average_dau_last_28_days,
                ROUND((SELECT AVG(behavior_count) FROM last_28), 2)
                    AS average_behavior_last_28_days
            FROM latest
        """).df()
        trend = connection.sql("""
            SELECT *
            FROM week2_anomaly_monitoring
            ORDER BY date_day
        """).df()
        anomalies = trend[trend["is_anomaly"] == 1].copy()
        segments = connection.sql("""
            SELECT
                user_segment,
                COUNT(*) AS user_count,
                ROUND(AVG(behavior_count), 2) AS avg_behavior_count,
                ROUND(AVG(content_count), 2) AS avg_content_count
            FROM week2_user_segments
            GROUP BY user_segment
            ORDER BY MIN(activity_quartile)
        """).df()
        retention = connection.sql("""
            SELECT *
            FROM week2_cohort_retention
            ORDER BY cohort_month, period_number
        """).df()
        retention_summary = (
            retention.groupby("period_number", as_index=False)
            .agg(
                cohort_count=("cohort_month", "nunique"),
                average_retention=("retention_rate", "mean"),
            )
        )
        content = connection.sql("""
            SELECT
                c.content_id,
                c.title,
                c.genres,
                COUNT(*) AS behavior_count,
                COUNT(DISTINCT b.user_id) AS active_users,
                ROUND(AVG(b.rating), 2) AS avg_rating
            FROM behavior_fact b
            JOIN content_dim c ON b.content_id = c.content_id
            GROUP BY c.content_id, c.title, c.genres
            HAVING COUNT(*) >= 100
            ORDER BY active_users DESC
            LIMIT 100
        """).df()
        concentration = connection.sql("""
            SELECT *
            FROM week2_content_concentration
            ORDER BY month_start
        """).df()
        cause_summary = (
            anomalies.groupby("diagnostic_reason", as_index=False)
            .agg(anomaly_count=("date_day", "count"))
            .sort_values("anomaly_count", ascending=False)
        )

        latest = overview.iloc[0]
        top_content = content.iloc[0]
        top_segment = segments.sort_values("user_count", ascending=False).iloc[0]
        top_concentration = concentration.loc[
            concentration["top10_behavior_share"].idxmax()
        ]
        d7 = retention_summary.loc[
            retention_summary["period_number"] == 1, "average_retention"
        ]
        average_period_1_retention = float(d7.iloc[0]) if not d7.empty else 0.0

        actions = pd.DataFrame(
            [
                {
                    "priority": "high",
                    "finding": "Activity segments differ in observed rating frequency.",
                    "action": "Use activity segment as a targeting variable for a recommendation or re-engagement test.",
                    "validation": "Compare seven-day rating-active rate by randomized user treatment.",
                },
                {
                    "priority": "high",
                    "finding": f"{top_content['title']} leads qualifying content by active users.",
                    "action": "Inspect related-content and recommendation exposure opportunities around high-reach titles.",
                    "validation": "Measure content diversity and seven-day activity, not only clicks or ratings.",
                },
                {
                    "priority": "medium",
                    "finding": f"Top-ten concentration peaks at {top_concentration['top10_behavior_share']:.2%} in {top_concentration['month_start']}.",
                    "action": "Monitor concentration to avoid over-reliance on a small content set.",
                    "validation": "Track long-tail active users and content coverage after recommendation changes.",
                },
                {
                    "priority": "high",
                    "finding": f"The anomaly monitor flagged {len(anomalies):,} dates.",
                    "action": "Route flagged dates through data-quality and product-event investigation before taking action.",
                    "validation": "Require pipeline completeness and cause confirmation for incident closure.",
                },
            ]
        )

        dashboard_data = {
            "overview": overview,
            "daily_trend": trend,
            "anomaly_dates": anomalies,
            "user_segments": segments,
            "retention": retention,
            "retention_summary": retention_summary,
            "content_ranking": content,
            "concentration": concentration,
            "cause_summary": cause_summary,
            "actions": actions,
        }
        output_path = OUTPUT_DIR / "ml_day14_dashboard.xlsx"
        with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
            for sheet_name, frame in dashboard_data.items():
                frame.to_excel(writer, sheet_name=sheet_name[:31], index=False)

        workbook = load_workbook(output_path)
        dashboard = workbook.create_sheet("dashboard", 0)
        dashboard.sheet_view.showGridLines = False
        dashboard["A1"] = "Simulated Content Platform Dashboard"
        dashboard["A1"].font = Font(size=18, bold=True, color="FFFFFF")
        dashboard["A1"].fill = PatternFill("solid", fgColor="1F4E78")
        dashboard.merge_cells("A1:H1")
        kpis = [
            ("Latest date", str(latest["latest_date"])),
            ("Latest DAU", int(latest["latest_dau"])),
            ("28-day average DAU", latest["average_dau_last_28_days"]),
            ("Potential anomalies", len(anomalies)),
        ]
        for column, (label, value) in enumerate(kpis, start=1):
            dashboard.cell(row=3, column=column, value=label)
            dashboard.cell(row=4, column=column, value=value)
            dashboard.cell(row=3, column=column).font = Font(bold=True)
            dashboard.cell(row=4, column=column).font = Font(size=14, bold=True)

        for sheet_name in dashboard_data:
            format_sheet(workbook[sheet_name[:31]])
        add_trend_chart(dashboard, workbook["daily_trend"])
        add_content_chart(dashboard, workbook["content_ranking"])
        add_segment_chart(dashboard, workbook["user_segments"])
        dashboard.column_dimensions["A"].width = 22
        for column in "BCDEFGHIJKLMNOP":
            dashboard.column_dimensions[column].width = 15
        workbook.save(output_path)

        report = f"""# Week 2 Analysis Report Draft

## 1. Business problem

This project simulates a content-platform analysis using MovieLens 25M. The business questions are why users may become inactive, which content attracts activity, and why core metrics fluctuate. Rating behavior is used as an activity proxy.

## 2. Metric definitions

- DAU: distinct users with at least one rating on a date.
- New-user proxy: users whose first observed rating occurs on the date.
- Content active users: distinct users rating a content item.
- Retention: users active in a later exact day or cohort month divided by the cohort size.
- Content concentration: the share of monthly rating behavior from the top ten content items.
- Anomaly: a date exceeding the configured historical z-score or change-rate threshold.

## 3. Data findings

- Latest observed date: {latest['latest_date']}.
- Latest DAU: {int(latest['latest_dau']):,}; 28-day average DAU: {latest['average_dau_last_28_days']:,.2f}.
- Highest qualifying content by active users: {top_content['title']} with {int(top_content['active_users']):,} active users.
- Largest user segment by count: {top_segment['user_segment']} with {int(top_segment['user_count']):,} users.
- Average month-1 cohort retention where available: {average_period_1_retention:.2%}.
- Highest top-ten content concentration: {top_concentration['top10_behavior_share']:.2%} in {top_concentration['month_start']}.
- Potential anomaly dates: {len(anomalies):,}.

## 4. Cause judgment

The results show associations in observed rating data, not causal explanations. The anomaly workbook compares behavior volume, active content, and new-user proxy changes to provide an initial diagnostic direction. Product logs, campaign logs, and pipeline logs would be required to confirm a cause.

## 5. Optimization suggestions

1. Use activity segments to target personalized recommendation or re-engagement treatments.
2. Inspect high-reach content and related-content paths while monitoring long-tail coverage.
3. Monitor concentration and retention together so that activity growth is not achieved by reducing content diversity.
4. Route statistical anomaly flags to a standard data-quality and business-event investigation process.

## 6. Validation plan

Run the Day 13 user-level A/B experiment plan. Use seven-day rating-active rate as the primary metric, pre-declare alpha = 0.05, report a 95% confidence interval, and monitor activity frequency, content diversity, rating quality, sample-ratio mismatch, and event completeness as guardrails.

## 7. Data limitations

MovieLens has ratings and tags but no registrations, impressions, clicks, watch time, user demographics, campaigns, recommendation exposure, or product-release logs. Cohorts, churn, retention, and anomalies are proxies based on observed ratings. This is a simulated public-data project and not a real company business result.

Generated on: {date.today().isoformat()}
"""
        (OUTPUT_DIR / "ml_day14_report_draft.md").write_text(
            report,
            encoding="utf-8",
        )
        print(f"Saved dashboard: {output_path}")
        print(f"Saved report: {OUTPUT_DIR / 'ml_day14_report_draft.md'}")
    finally:
        connection.close()


if __name__ == "__main__":
    main()
