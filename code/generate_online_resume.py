from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


PROJECT_DIR = Path(__file__).resolve().parents[1]
RESUME_DIR = PROJECT_DIR / "简历"
OUTPUT_DIR = PROJECT_DIR / "days"
PHOTO_FILE = RESUME_DIR / "证件照_简历版.jpg"
OUTPUT_FILE = RESUME_DIR / "张博能_数据分析实习_网申版.docx"
MARKDOWN_FILE = RESUME_DIR / "张博能_数据分析实习_网申版.md"

NAME = "张博能"
PHONE = "18974648806"
EMAIL = "1786013324@qq.com"
LOCATION = "湘潭"


def get_metrics():
    forest_auc, forest_f1 = "0.860", "0.719"
    model_file = OUTPUT_DIR / "ml_day18_model_comparison.csv"
    if model_file.exists():
        data = pd.read_csv(model_file)
        row = data.loc[
            data["model"] == "random_forest_with_enhanced_features"
        ].iloc[0]
        forest_auc = f"{row['roc_auc']:.3f}"
        forest_f1 = f"{row['f1']:.3f}"
    return forest_auc, forest_f1


FOREST_AUC, FOREST_F1 = get_metrics()


def set_font(run, size=10, bold=False, color=(64, 64, 64)):
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor(*color)


def remove_table_borders(table):
    properties = table._tbl.tblPr
    borders = properties.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "nil")


def add_heading(document, text):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(7)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.line_spacing = 1.0
    run = paragraph.add_run(text)
    set_font(run, size=12, bold=True, color=(31, 78, 121))
    paragraph_border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "3")
    bottom.set(qn("w:color"), "1F4E79")
    paragraph_border.append(bottom)
    paragraph._p.get_or_add_pPr().append(paragraph_border)


def add_bullet(document, text):
    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.left_indent = Cm(0.42)
    paragraph.paragraph_format.first_line_indent = Cm(-0.18)
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.line_spacing = 1.03
    run = paragraph.add_run(text)
    set_font(run, size=9.6)


def add_line(document, text, size=9.8, bold=False, after=2):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1.05
    run = paragraph.add_run(text)
    set_font(run, size=size, bold=bold)


def build_document():
    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(1.0)
    section.bottom_margin = Cm(0.8)
    section.left_margin = Cm(1.35)
    section.right_margin = Cm(1.35)

    normal = document.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(9.6)

    header = document.add_table(rows=1, cols=2)
    header.autofit = False
    header.columns[0].width = Cm(14.8)
    header.columns[1].width = Cm(2.8)
    remove_table_borders(header)
    left, right = header.rows[0].cells

    title = left.paragraphs[0]
    title.paragraph_format.space_after = Pt(2)
    run = title.add_run(NAME)
    set_font(run, size=24, bold=True, color=(31, 78, 121))

    target = left.add_paragraph()
    target.paragraph_format.space_after = Pt(4)
    run = target.add_run("求职方向：数据分析 / 数据分析实习")
    set_font(run, size=11.5, bold=True, color=(80, 80, 80))

    contact = left.add_paragraph()
    contact.paragraph_format.space_after = Pt(1)
    run = contact.add_run(
        f"男｜现居湘潭｜{PHONE}｜{EMAIL}"
    )
    set_font(run, size=9.5)

    photo = right.paragraphs[0]
    photo.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if PHOTO_FILE.exists():
        photo.add_run().add_picture(
            str(PHOTO_FILE),
            width=Cm(2.65),
            height=Cm(3.4),
        )

    add_heading(document, "教育背景")
    add_line(
        document,
        "2028.06 预计毕业｜湖南科技大学｜计算机科学与工程｜本科在读",
        size=10,
        bold=True,
    )
    add_line(document, "英语：大学英语六级（CET-6）504分", size=9.5, after=3)

    add_heading(document, "项目经历")
    add_line(
        document,
        "泛文娱内容平台增长与推荐分析｜MovieLens 25M公开数据模拟项目",
        size=10.2,
        bold=True,
        after=3,
    )
    add_bullet(
        document,
        "面对2500万条评分数据，设计用户—内容—行为—日期四表模型，用 DuckDB 完成留存、漏斗、Top N、窗口函数和异常检测等20+项分析查询，形成日度指标监控和用户分群结果。",
    )
    add_bullet(
        document,
        "用 pandas 完成数据质量检查和 EDA，发现用户活跃度呈长尾分布；高活跃分群平均行为数429.78，低活跃分群26.43，前者约为后者16.3倍，据此提出分群运营建议。",
    )
    add_bullet(
        document,
        f"按时间切分构造未来30天活跃预测任务，对比三个模型；Random Forest 测试 ROC-AUC {FOREST_AUC}、F1 {FOREST_F1}，并完成特征工程、数据泄漏审计和误判样本分析。",
    )
    add_bullet(
        document,
        "用 PySpark 完成抽样数据的 groupBy、join、窗口排名、分区和物理执行计划分析；用 TF-IDF + 逻辑回归完成标签文本分类基线。",
    )

    add_heading(document, "专业技能")
    add_bullet(
        document,
        "熟练：SQL（多表关联、CTE、窗口函数、留存、漏斗、Top N、异常检测）；Python（pandas、NumPy、EDA、可视化）。",
    )
    add_bullet(
        document,
        "掌握：scikit-learn（逻辑回归、随机森林、时间切分、模型评估）；Excel（数据透视、指标看板、条件格式）。",
    )
    add_bullet(
        document,
        "了解：PyTorch（MLP训练循环）；PySpark（DataFrame、groupBy、join、窗口函数）；Transformer基本概念与文本分类流程。",
    )

    add_heading(document, "个人优势")
    add_line(
        document,
        "完成1个端到端公开数据项目，能够独立完成 SQL 建模、Python 分析、Excel 看板和 A/B 实验设计；已完成20+项项目SQL分析查询。对公开数据中的代理指标和数据局限保持诚实说明。",
        size=9.6,
        after=2,
    )
    return document


def main():
    RESUME_DIR.mkdir(parents=True, exist_ok=True)
    document = build_document()
    document.save(OUTPUT_FILE)
    markdown = """# 张博能｜数据分析 / 数据分析实习

简历 Word 文件：`张博能_数据分析实习_网申版.docx`

该版本为线上投递单栏简历，聚焦数据分析方向。未加入尚未验证的 GitHub、100+ SQL 题和 90 分钟限时测试信息。
"""
    MARKDOWN_FILE.write_text(markdown, encoding="utf-8")
    print(f"Saved: {OUTPUT_FILE}")


if __name__ == "__main__":
    main()
