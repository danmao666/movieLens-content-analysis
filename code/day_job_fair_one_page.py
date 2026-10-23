from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from common import OUTPUT_DIR


PROJECT_DIR = Path(__file__).resolve().parents[1]
MATCH_FILE = OUTPUT_DIR / "ml_job_match_analysis_updated.xlsx"
GUIDE_FILE = OUTPUT_DIR / "ml_job_fair_tomorrow_guide.xlsx"
OUTPUT_FILE = OUTPUT_DIR / "ml_job_fair_tomorrow_guide_updated.xlsx"


TECH_MATCH = {
    "杭州卡日曲网络科技有限公司": (
        "SQL、Python/pandas、DAU/留存/cohort、用户分层、A/B实验、内容推荐分析",
        "游戏次留/付费/关卡漏斗、活动运营、游戏数值和实时运营平台",
    ),
    "湖南向上商旅服务有限公司": (
        "SQL、Python/pandas、Excel、数据清洗、指标计算、可视化、异常分析",
        "机票/商旅业务指标、生产BI/ETL、业务数据口径；逆向爬虫不作为主投能力",
    ),
    "中国人民财产保险股份有限公司湘潭市分公司": (
        "SQL、Python、统计检验、用户/风险分层、异常监控、基础预测模型",
        "保险产品、车险理赔、定价/风控、精算和保险数据口径",
    ),
    "中国平安财产保险股份有限公司湖北分公司": (
        "Python/SQL、指标体系、Excel看板、统计分析、模型基础、业务报告",
        "保险经营指标、数据治理、金融业务和管培岗位具体方向；需确认27届资格",
    ),
    "意拉德电子（东莞）有限公司": (
        "Python、SQL、pandas、机器学习、PyTorch基础、PySpark基础、需求拆解和报告",
        "企业应用/API、系统实施、部署、接口调试、生产级AI工程；实习身份需确认",
    ),
    "广州英码信息科技有限公司": (
        "SQL、Python、数据分析、项目复盘、需求整理、技术表达和Excel看板",
        "C/C++、硬件/驱动、软件交付、部署、自动化测试和客户现场经验",
    ),
    "杰创智能科技股份有限公司": (
        "指标分析、Python/SQL、看板、数据解释、报告和跨团队沟通",
        "数据中心运维、网络、嵌入式、前端、机器人算法和智能硬件开发",
    ),
    "中科开创（广州）智能科技发展有限公司": (
        "数据分析、指标解释、项目讲解、Python/SQL和结构化排查",
        "硬件/设备、行业解决方案、售前技术营销和客户现场支持",
    ),
    "浙江省公众信息产业有限公司": (
        "内容平台分析、用户分层、留存、A/B方案、Python/SQL和运营看板",
        "AI运营流程、提示词/知识库、LLM应用、客户运营和大数据平台",
    ),
    "杭州铭师堂数字科技有限公司": (
        "业务问题拆解、指标定义、用户分析、看板、实验方案和项目表达",
        "PRD、用户研究、产品生命周期、研发协作和互联网产品实战",
    ),
    "乔锋智能装备股份有限公司": (
        "SQL、Excel、数据核对、Python数据处理、指标分析和运营报告",
        "ERP/OA、桌面/网络支持、制造业运营和企业IT故障处理",
    ),
    "浙江永太科技股份有限公司": (
        "Python/SQL、数据处理、报表和基础问题排查",
        "Linux/网络、权限、桌面运维、硬件、工控和自控系统；优先问IT专员职责",
    ),
    "巴德富集团有限公司": (
        "Python、SQL、基础机器学习、PyTorch基础、指标分析和应用验证",
        "企业应用/API、权限、部署、Linux、安全和生产级AI应用",
    ),
    "长沙迪迈科技股份有限公司": (
        "AI/数据基础、指标体系、业务分析、项目报告和解决方案表达",
        "矿业领域、AI产品设计、客户访谈、咨询规划和售前交付",
    ),
    "广东德晟智能科技有限公司": (
        "SQL、数据质量检查、异常排查、项目管理基础和分析报告",
        "制造质量/测试流程、自动化测试、硬件/工程知识和NPI流程",
    ),
    "深圳出海一叮科技有限公司长沙分公司": (
        "用户分层、指标分析、A/B实验、内容运营和Excel/Python分析",
        "SaaS产品生命周期、跨境电商、广告投放、海外平台和英文业务",
    ),
    "湖南长银五八消费金融股份有限公司": (
        "统计、基础模型、异常检测、用户分层和指标监控",
        "金融风控、评分卡、授信、逾期/坏账、监管和风险策略",
    ),
    "湖南硅谷融创科技有限公司": (
        "SQL、Python、数据核对、报表、问题拆解和项目文档",
        "Linux、网络、数据库运维、部署、日志和鸿蒙测试；开发岗不匹配",
    ),
    "湖南亚太实业有限公司": (
        "SQL、Excel、Python数据处理、业务流程分析和报表",
        "ERP/OA、权限管理、网络、企业信息化实施和系统维护",
    ),
    "拓维信息系统股份有限公司": (
        "内容平台分析、推荐策略、用户留存、A/B实验、Python/SQL和基础模型",
        "LLM产品落地、AI解决方案交付、游戏产品运营和软件开发；需确认27届",
    ),
}


def format_sheet(path: Path) -> None:
    workbook = load_workbook(path)
    sheet = workbook["明日现场清单"]
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A7"
    sheet.auto_filter.ref = f"A6:I{sheet.max_row}"
    sheet.merge_cells("A1:I1")
    sheet.merge_cells("A2:I2")
    sheet.merge_cells("B3:I3")
    sheet.merge_cells("B4:I4")
    sheet["A1"].font = Font(size=18, bold=True, color="FFFFFF")
    sheet["A1"].fill = PatternFill("solid", fgColor="1F4E78")
    sheet["A1"].alignment = Alignment(horizontal="center", vertical="center")
    sheet["A2"].font = Font(bold=True, color="404040")
    sheet["A2"].alignment = Alignment(horizontal="center", vertical="center")
    for cell_ref in ("A3", "A4"):
        sheet[cell_ref].font = Font(bold=True, color="FFFFFF")
        sheet[cell_ref].fill = PatternFill("solid", fgColor="70AD47")
        sheet[cell_ref].alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )
    for cell_ref in ("B3", "B4"):
        sheet[cell_ref].alignment = Alignment(
            vertical="top",
            wrap_text=True,
        )
        sheet[cell_ref].fill = PatternFill("solid", fgColor="E2F0D9")
    sheet.row_dimensions[1].height = 32
    sheet.row_dimensions[2].height = 24
    sheet.row_dimensions[3].height = 68
    sheet.row_dimensions[4].height = 68

    header_fill = PatternFill("solid", fgColor="5B9BD5")
    border = Border(bottom=Side(style="thin", color="D9E2F3"))
    for cell in sheet[6]:
        cell.fill = header_fill
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )
        cell.border = border
    sheet.row_dimensions[6].height = 30

    widths = {
        "A": 10,
        "B": 10,
        "C": 28,
        "D": 48,
        "E": 18,
        "F": 10,
        "G": 54,
        "H": 54,
        "I": 58,
    }
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width

    for row in sheet.iter_rows(min_row=6):
        max_lines = 1
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = border
            value = "" if cell.value is None else str(cell.value)
            width = sheet.column_dimensions[cell.column_letter].width or 12
            max_lines = max(max_lines, (len(value) // max(int(width) - 2, 1)) + 1)
        sheet.row_dimensions[row[0].row].height = min(
            max(46, max_lines * 15),
            150,
        )

    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_title_rows = "1:6"
    workbook.save(path)


def main() -> None:
    if not MATCH_FILE.exists():
        raise FileNotFoundError(f"Missing match workbook: {MATCH_FILE}")
    if not GUIDE_FILE.exists():
        raise FileNotFoundError(f"Missing booth guide: {GUIDE_FILE}")

    direct = pd.read_excel(MATCH_FILE, sheet_name="A_direct")
    conditional = pd.read_excel(MATCH_FILE, sheet_name="B_conditional")
    matches = pd.concat([direct, conditional], ignore_index=True)
    matches = matches.drop_duplicates(subset=["company", "role"])

    booths = pd.read_excel(GUIDE_FILE, sheet_name="按展位号")
    booths = booths.drop_duplicates(subset=["company"])
    booth_columns = booths[[
        "company",
        "booth",
        "company_city",
        "headcount",
    ]].rename(columns={"headcount": "booth_headcount"})
    result = matches.merge(booth_columns, on="company", how="inner")
    result["priority"] = result["level"].astype(str).str.startswith("A").map(
        {True: "A优先", False: "B可投"}
    )
    def join_unique(values):
        return "；".join(dict.fromkeys(str(value) for value in values))

    result = (
        result.groupby(
            [
                "priority",
                "company",
                "booth",
                "company_city",
                "booth_headcount",
            ],
            as_index=False,
        )
        .agg(
            role=("role", join_unique),
            level=("level", "first"),
            reason=("reason", "first"),
            gap=("gap", join_unique),
        )
    )
    result["matched_tech"] = result["company"].map(
        lambda company: TECH_MATCH.get(
            company,
            ("SQL、Python、Excel、项目分析", "需现场确认岗位具体技术要求"),
        )[0]
    )
    result["missing_tech"] = result["company"].map(
        lambda company: TECH_MATCH.get(
            company,
            ("SQL、Python、Excel、项目分析", "需现场确认岗位具体技术要求"),
        )[1]
    )
    result["on_site_question"] = result.apply(
        lambda row: (
            f"确认“{row['role']}”中的具体岗位；说明SQL/Python/Excel项目；"
            "询问专业、届次、笔试和面试流程。"
            + f" 重点补问：{row['missing_tech']}"
        ),
        axis=1,
    )
    result = result.sort_values(
        ["priority", "booth", "company"],
        key=lambda column: column.map({"A优先": 0, "B可投": 1})
        if column.name == "priority"
        else column,
    )
    result = result[[
        "priority",
        "booth",
        "company",
        "role",
        "company_city",
        "booth_headcount",
        "matched_tech",
        "missing_tech",
        "on_site_question",
    ]].rename(
        columns={
            "priority": "优先级",
            "booth": "展位号",
            "company": "单位名称",
            "role": "建议岗位",
            "company_city": "公司所在地",
            "booth_headcount": "需求人数",
            "matched_tech": "我已满足的技术",
            "missing_tech": "不足/需补技术",
            "on_site_question": "现场重点",
        }
    )

    profile = pd.read_excel(MATCH_FILE, sheet_name="profile")
    engineering_gap = pd.read_excel(
        MATCH_FILE,
        sheet_name="C_core_engineering_gap",
    )
    strong = profile.loc[
        profile["section"] == "可独立展示", "skills"
    ].iloc[0]
    basic = profile.loc[
        profile["section"] == "基础接触", "skills"
    ].iloc[0]
    missing = profile.loc[
        profile["section"] == "当前缺口", "skills"
    ].iloc[0]
    gap_roles = "；".join(engineering_gap["roles"].astype(str).tolist())
    matched_technology = f"{strong}；基础接触：{basic}"
    unmatched_technology = f"{missing}；对应岗位类型：{gap_roles}"

    with pd.ExcelWriter(OUTPUT_FILE, engine="openpyxl") as writer:
        result.to_excel(writer, sheet_name="明日现场清单", startrow=5, index=False)
    workbook = load_workbook(OUTPUT_FILE)
    sheet = workbook["明日现场清单"]
    sheet["A1"] = "明日招聘会现场清单"
    sheet["A2"] = (
        "湖南科技大学月湖广场｜2026-10-17 9:00-12:00｜按展位号找单位，C类核心工程岗位不列入"
    )
    sheet["A3"] = "我能匹配的技术"
    sheet["B3"] = matched_technology
    sheet["A4"] = "当前不匹配/需补强"
    sheet["B4"] = unmatched_technology
    workbook.save(OUTPUT_FILE)
    format_sheet(OUTPUT_FILE)
    result.to_csv(
        OUTPUT_DIR / "ml_job_fair_one_page_booth_sheet.csv",
        index=False,
        encoding="utf-8-sig",
    )
    print(f"Saved: {OUTPUT_FILE}")
    print(f"Rows: {len(result)}")


if __name__ == "__main__":
    main()
