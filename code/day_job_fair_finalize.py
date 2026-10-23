from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from common import OUTPUT_DIR
from day_job_fair_one_page import TECH_MATCH


GUIDE_FILE = OUTPUT_DIR / "ml_job_fair_tomorrow_guide.xlsx"
MATCH_FILE = OUTPUT_DIR / "ml_job_match_analysis_updated.xlsx"
OUTPUT_FILE = GUIDE_FILE
FALLBACK_OUTPUT = OUTPUT_DIR / "ml_job_fair_tomorrow_guide_final.xlsx"


def format_workbook(path: Path) -> None:
    workbook = load_workbook(path)
    sheet = workbook["明日现场清单"]
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:K{sheet.max_row}"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True

    header_fill = PatternFill("solid", fgColor="1F4E78")
    border = Border(bottom=Side(style="thin", color="D9E2F3"))
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )
        cell.border = border
    sheet.row_dimensions[1].height = 34

    widths = {
        "A": 10,
        "B": 10,
        "C": 28,
        "D": 48,
        "E": 18,
        "F": 10,
        "G": 20,
        "H": 54,
        "I": 54,
        "J": 58,
        "K": 70,
    }
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width

    for row in sheet.iter_rows(min_row=2):
        max_lines = 1
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = border
            value = "" if cell.value is None else str(cell.value)
            width = sheet.column_dimensions[cell.column_letter].width or 12
            max_lines = max(
                max_lines,
                (len(value) // max(int(width) - 2, 1)) + 1,
            )
        sheet.row_dimensions[row[0].row].height = min(
            max(46, max_lines * 15),
            150,
        )
    sheet.print_title_rows = "1:1"
    workbook.save(path)


def load_match_fallback():
    direct = pd.read_excel(MATCH_FILE, sheet_name="A_direct")
    conditional = pd.read_excel(MATCH_FILE, sheet_name="B_conditional")
    matches = pd.concat([direct, conditional], ignore_index=True)
    return (
        matches.groupby("company", as_index=False)
        .agg(
            reason=("reason", "first"),
            gap=("gap", lambda values: "；".join(dict.fromkeys(values))),
        )
    )


def main() -> None:
    if not GUIDE_FILE.exists():
        raise FileNotFoundError(f"Missing guide workbook: {GUIDE_FILE}")
    if not MATCH_FILE.exists():
        raise FileNotFoundError(f"Missing match workbook: {MATCH_FILE}")

    guide = pd.read_excel(GUIDE_FILE, sheet_name="按展位号")
    fallback = load_match_fallback()
    guide = guide.merge(fallback, on="company", how="left")

    def tech_pair(company):
        return TECH_MATCH.get(
            company,
            (
                "根据项目具备SQL、Python、Excel和基础数据分析能力",
                "需现场确认该岗位的具体技术要求",
            ),
        )

    guide[["matched_tech", "missing_tech"]] = guide["company"].apply(
        lambda company: pd.Series(tech_pair(company))
    )
    guide["site_focus"] = guide.apply(
        lambda row: (
            f"确认具体岗位：{row['recommended_roles']}；"
            "说明已完成的SQL/Python/Excel项目；"
            "询问专业、届次、笔试和面试流程；"
            f"重点确认：{row['missing_tech']}"
        ),
        axis=1,
    )

    final = guide[[
        "priority",
        "booth",
        "company",
        "recommended_roles",
        "company_city",
        "headcount",
        "source",
        "matched_tech",
        "missing_tech",
        "site_focus",
        "all_listed_positions",
    ]].rename(
        columns={
            "priority": "优先级",
            "booth": "展位号",
            "company": "单位名称",
            "recommended_roles": "建议岗位",
            "company_city": "公司所在地",
            "headcount": "需求人数",
            "source": "来源表",
            "matched_tech": "匹配的技术",
            "missing_tech": "欠缺/需补技术",
            "site_focus": "现场重点",
            "all_listed_positions": "单位全部岗位",
        }
    )
    final = final.sort_values(
        ["优先级", "展位号"],
        key=lambda column: column.map({"A": 0, "B": 1})
        if column.name == "优先级"
        else column,
    )

    output = OUTPUT_FILE
    try:
        with pd.ExcelWriter(output, engine="openpyxl") as writer:
            final.to_excel(writer, sheet_name="明日现场清单", index=False)
    except PermissionError:
        output = FALLBACK_OUTPUT
        with pd.ExcelWriter(output, engine="openpyxl") as writer:
            final.to_excel(writer, sheet_name="明日现场清单", index=False)
        print("Original guide is open or locked; saved a final copy instead.")

    format_workbook(output)
    final.to_csv(
        OUTPUT_DIR / "ml_job_fair_tomorrow_guide_final.csv",
        index=False,
        encoding="utf-8-sig",
    )
    print(f"Saved: {output}")
    print(f"Rows: {len(final)}")


if __name__ == "__main__":
    main()
