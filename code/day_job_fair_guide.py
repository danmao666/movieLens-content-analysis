from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill, Side, Border

from common import OUTPUT_DIR


DOWNLOAD_DIR = Path.home() / "Downloads"
FILES = [
    (
        "不限专业表",
        DOWNLOAD_DIR
        / "湖南科技大学2026年秋季毕业生供需见面会参会单位岗位信息-不限专业.xlsx",
    ),
    (
        "计算机相关专业表",
        DOWNLOAD_DIR
        / "湖南科技大学2026年秋季毕业生供需见面会参会单位岗位信息-计算机相关专业.xlsx",
    ),
]


TARGETS = [
    ("A", "湖南向上商旅服务有限公司", "数据分析员"),
    ("A", "意拉德电子（东莞）有限公司", "大数据及人工智能应用实习生；应用系统需求分析及方案实施实习生"),
    ("A", "中国人民财产保险股份有限公司湘潭市分公司", "车险数据分析岗"),
    ("A", "中国平安财产保险股份有限公司湖北分公司", "数据管培生"),
    ("A", "杭州卡日曲网络科技有限公司", "游戏增长策略管培生；游戏运营管培生"),
    ("B", "广州英码信息科技有限公司", "产品助理；交付支持工程师；软件交付工程师"),
    ("B", "杰创智能科技股份有限公司", "技术支持工程师；销售运营支持"),
    ("B", "中科开创（广州）智能科技发展有限公司", "技术支持工程师；技术营销工程师"),
    ("B", "浙江省公众信息产业有限公司", "AI运营工程师；互联网产品运营"),
    ("B", "杭州铭师堂数字科技有限公司", "互联网类产品专员"),
    ("B", "乔锋智能装备股份有限公司", "IT助理；运营管培生"),
    ("B", "浙江永太科技股份有限公司", "IT专员"),
    ("B", "巴德富集团有限公司", "AI应用助理工程师；IT应用助理工程师"),
    ("B", "长沙迪迈科技股份有限公司", "AI产品经理；解决方案工程师；咨询规划师"),
    ("B", "广东德晟智能科技有限公司", "技术支持工程师；项目工程师"),
    ("B", "深圳出海一叮科技有限公司长沙分公司", "SaaS产品经理；海外市场运营"),
    ("B", "湖南长银五八消费金融股份有限公司", "全面风险管理岗"),
    ("B", "湖南硅谷融创科技有限公司", "实施运维工程师"),
    ("B", "湖南亚太实业有限公司", "信息化助理工程师"),
    ("B", "拓维信息系统股份有限公司", "AI产品经理；AI解决方案；游戏产品运营"),
]


def load_jobs():
    frames = []
    for source, path in FILES:
        if not path.exists():
            raise FileNotFoundError(f"岗位文件不存在: {path}")
        frame = pd.read_excel(path)
        frame["source"] = source
        frames.append(frame)
    return pd.concat(frames, ignore_index=True)


def format_workbook(path):
    workbook = load_workbook(path)
    widths = {
        "priority": 12,
        "booth": 12,
        "company": 30,
        "recommended_roles": 58,
        "company_city": 18,
        "headcount": 12,
        "source": 20,
        "venue": 28,
        "time": 24,
        "all_listed_positions": 70,
    }
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.sheet_view.showGridLines = False
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )
        sheet.row_dimensions[1].height = 30
        for cell in sheet[1]:
            sheet.column_dimensions[cell.column_letter].width = widths.get(
                str(cell.value), 18
            )
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                cell.border = Border(
                    bottom=Side(style="thin", color="D9E2F3")
                )
            sheet.row_dimensions[row[0].row].height = 54
    workbook.save(path)


def main():
    source = load_jobs()
    records = []
    for priority, company, roles in TARGETS:
        matches = source[
            source["单位名称"].astype(str).str.strip() == company
        ]
        if matches.empty:
            continue
        match = matches.sort_values("展位号").iloc[0]
        records.append(
            {
                "priority": priority,
                "booth": int(match["展位号"]),
                "company": company,
                "recommended_roles": roles,
                "company_city": f"{match['所在省份']}{match['所在城市']}",
                "headcount": match["需求人数"],
                "source": "；".join(sorted(matches["source"].unique())),
                "venue": match["双选会地点"],
                "time": match["举办时间"],
                "all_listed_positions": match["招聘岗位"],
            }
        )

    guide = pd.DataFrame(records).sort_values(
        ["booth", "priority", "company"]
    )
    priority_view = guide.sort_values(
        ["priority", "booth", "company"]
    )
    notes = pd.DataFrame(
        [
            {
                "item": "现场地点",
                "value": "湖南科技大学月湖广场",
                "note": "表中的“双选会地点”，不是公司所在地。",
            },
            {
                "item": "现场时间",
                "value": "2026-10-17 9:00-12:00",
                "note": "建议提前到场，先按展位号找A档单位。",
            },
            {
                "item": "找公司方法",
                "value": "看现场展位号牌",
                "note": "公司所在地如广州、长沙、杭州不代表公司在现场的地理位置。",
            },
            {
                "item": "沟通重点",
                "value": "先问具体岗位是否接受计算机/数据分析背景",
                "note": "一行可能包含多个岗位，现场要确认具体岗位和专业限制。",
            },
        ]
    )

    output = OUTPUT_DIR / "ml_job_fair_tomorrow_guide.xlsx"
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        guide.to_excel(writer, sheet_name="按展位号", index=False)
        priority_view.to_excel(writer, sheet_name="按优先级", index=False)
        notes.to_excel(writer, sheet_name="现场说明", index=False)
    format_workbook(output)
    guide.to_csv(
        OUTPUT_DIR / "ml_job_fair_tomorrow_guide.csv",
        index=False,
        encoding="utf-8-sig",
    )
    print(f"Saved: {output}")


if __name__ == "__main__":
    main()
