import pandas as pd
from openpyxl import load_workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.formatting.rule import CellIsRule, DataBarRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from common import OUTPUT_DIR, connect, require_tables


def format_data_sheet(worksheet):
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.sheet_view.showGridLines = False

    for cell in worksheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center")

    for column_cells in worksheet.columns:
        column_letter = get_column_letter(column_cells[0].column)
        max_length = max(
            len(str(cell.value)) if cell.value is not None else 0
            for cell in column_cells
        )
        worksheet.column_dimensions[column_letter].width = min(
            max(max_length + 2, 12), 40
        )


def add_daily_chart(dashboard, daily_sheet):
    last_row = daily_sheet.max_row
    chart = LineChart()
    chart.title = "Daily Active Users"
    chart.y_axis.title = "DAU"
    chart.x_axis.title = "Date"
    chart.height = 8
    chart.width = 16

    categories = Reference(
        daily_sheet,
        min_col=1,
        min_row=2,
        max_row=last_row,
    )

    for column in (2, 7):
        data = Reference(
            daily_sheet,
            min_col=column,
            max_col=column,
            min_row=1,
            max_row=last_row,
        )
        chart.add_data(data, titles_from_data=True, from_rows=False)
    chart.set_categories(categories)
    chart.series[0].graphicalProperties.line.solidFill = "2F6B9A"
    if len(chart.series) > 1:
        chart.series[1].graphicalProperties.line.solidFill = "F28E2B"
    dashboard.add_chart(chart, "A8")


def add_content_chart(dashboard, content_sheet):
    last_row = min(content_sheet.max_row, 21)
    chart = BarChart()
    chart.type = "bar"
    chart.style = 10
    chart.title = "Top Content by Active Users"
    chart.x_axis.title = "Active Users"
    chart.y_axis.title = "Content"
    chart.height = 8
    chart.width = 16

    data = Reference(content_sheet, min_col=4, min_row=1, max_row=last_row)
    categories = Reference(content_sheet, min_col=1, min_row=2, max_row=last_row)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(categories)
    dashboard.add_chart(chart, "J8")


def add_segment_chart(dashboard, segment_sheet):
    chart = PieChart()
    chart.title = "User Segments"
    chart.height = 8
    chart.width = 12

    data = Reference(
        segment_sheet,
        min_col=2,
        min_row=1,
        max_row=segment_sheet.max_row,
    )
    labels = Reference(
        segment_sheet,
        min_col=1,
        min_row=2,
        max_row=segment_sheet.max_row,
    )
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(labels)
    dashboard.add_chart(chart, "A25")


def main() -> None:
    connection = connect()
    try:
        require_tables(connection, "daily_metrics", "behavior_fact", "content_dim")

        daily = connection.sql("""
            SELECT *
            FROM daily_metrics
            ORDER BY date_day
        """).df()

        content = connection.sql("""
            SELECT
                c.title,
                c.genres,
                COUNT(*) AS behavior_count,
                COUNT(DISTINCT b.user_id) AS active_users,
                ROUND(AVG(b.rating), 2) AS avg_rating
            FROM behavior_fact b
            JOIN content_dim c ON b.content_id = c.content_id
            GROUP BY c.content_id, c.title, c.genres
            HAVING COUNT(*) >= 100
            ORDER BY active_users DESC
            LIMIT 50
        """).df()

        segments = connection.sql("""
            SELECT
                CASE
                    WHEN behavior_count >= 50 THEN 'high'
                    WHEN behavior_count >= 20 THEN 'middle'
                    ELSE 'low'
                END AS user_segment,
                COUNT(*) AS user_count,
                ROUND(AVG(behavior_count), 2) AS avg_behavior_count
            FROM user_dim
            GROUP BY user_segment
            ORDER BY CASE user_segment
                WHEN 'high' THEN 1
                WHEN 'middle' THEN 2
                ELSE 3
            END
        """).df()

        daily["date_day"] = pd.to_datetime(daily["date_day"])
        pivot_source = daily.assign(
            month=daily["date_day"].dt.to_period("M").astype(str)
        )
        pivot_summary = pd.pivot_table(
            pivot_source,
            index="month",
            values=["dau", "behavior_count", "active_content_count"],
            aggfunc={
                "dau": ["sum", "mean"],
                "behavior_count": "sum",
                "active_content_count": "mean",
            },
        ).reset_index()
        pivot_summary.columns = [
            "_".join(str(part) for part in column if part).rstrip("_")
            if isinstance(column, tuple)
            else str(column)
            for column in pivot_summary.columns
        ]
        output_path = OUTPUT_DIR / "ml_day05_content_dashboard.xlsx"

        with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
            daily.to_excel(writer, sheet_name="daily_metrics", index=False)
            content.to_excel(writer, sheet_name="content_ranking", index=False)
            segments.to_excel(writer, sheet_name="user_segments", index=False)
            pivot_summary.to_excel(writer, sheet_name="pivot_summary", index=False)

            formula_dates = daily[["date_day"]].head(10).copy()
            formula_dates["xlookup_dau"] = ""
            formula_dates["sumifs_behavior_count"] = ""
            formula_dates.to_excel(writer, sheet_name="excel_practice", index=False)

        workbook = load_workbook(output_path)
        dashboard = workbook.create_sheet("dashboard", 0)
        dashboard.sheet_view.showGridLines = False
        dashboard["A1"] = "Content Platform Daily Dashboard"
        dashboard["A1"].font = Font(size=18, bold=True, color="FFFFFF")
        dashboard["A1"].fill = PatternFill("solid", fgColor="1F4E78")
        dashboard.merge_cells("A1:H1")

        latest = daily.iloc[-1]
        kpis = [
            ("Latest date", str(latest["date_day"].date())),
            ("Latest DAU", int(latest["dau"])),
            ("Behavior count", int(latest["behavior_count"])),
            ("Active content", int(latest["active_content_count"])),
        ]
        for index, (label, value) in enumerate(kpis, start=3):
            dashboard.cell(row=3, column=index, value=label)
            dashboard.cell(row=4, column=index, value=value)
            dashboard.cell(row=3, column=index).font = Font(bold=True)
            dashboard.cell(row=4, column=index).font = Font(size=14, bold=True)

        daily_sheet = workbook["daily_metrics"]
        content_sheet = workbook["content_ranking"]
        segment_sheet = workbook["user_segments"]
        practice_sheet = workbook["excel_practice"]

        for worksheet in (
            daily_sheet,
            content_sheet,
            segment_sheet,
            workbook["pivot_summary"],
            practice_sheet,
        ):
            format_data_sheet(worksheet)

        practice_sheet["A1"] = "date_day"
        practice_sheet["B1"] = "xlookup_dau"
        practice_sheet["C1"] = "sumifs_behavior_count"
        for row in range(2, practice_sheet.max_row + 1):
            practice_sheet.cell(
                row=row,
                column=2,
                value=(
                    f'=XLOOKUP(A{row},daily_metrics!$A:$A,'
                    f'daily_metrics!$B:$B,"not found")'
                ),
            )
            practice_sheet.cell(
                row=row,
                column=3,
                value=(
                    f'=SUMIFS(daily_metrics!$C:$C,'
                    f'daily_metrics!$A:$A,A{row})'
                ),
            )
        format_data_sheet(practice_sheet)

        daily_headers = {
            cell.value: cell.column for cell in daily_sheet[1]
        }
        anomaly_column = daily_headers.get("is_dau_anomaly")
        if anomaly_column:
            cells = f"{get_column_letter(anomaly_column)}2:{get_column_letter(anomaly_column)}{daily_sheet.max_row}"
            daily_sheet.conditional_formatting.add(
                cells,
                CellIsRule(
                    operator="equal",
                    formula=["1"],
                    fill=PatternFill("solid", fgColor="F4CCCC"),
                ),
            )

        dau_column_letter = get_column_letter(daily_headers["dau"])
        daily_sheet.conditional_formatting.add(
            f"{dau_column_letter}2:{dau_column_letter}{daily_sheet.max_row}",
            DataBarRule(
                start_type="min",
                end_type="max",
                color="5B9BD5",
                showValue=True,
            ),
        )

        add_daily_chart(dashboard, daily_sheet)
        add_content_chart(dashboard, content_sheet)
        add_segment_chart(dashboard, segment_sheet)
        dashboard.column_dimensions["A"].width = 18
        for column in "BCDEFGHIJKLMNOP":
            dashboard.column_dimensions[column].width = 14

        workbook.save(output_path)
        daily.to_csv(
            OUTPUT_DIR / "ml_day05_daily_metrics.csv",
            index=False,
            encoding="utf-8-sig",
        )
        content.to_csv(
            OUTPUT_DIR / "ml_day05_content_metrics.csv",
            index=False,
            encoding="utf-8-sig",
        )
        print(f"Saved: {output_path}")
    finally:
        connection.close()


if __name__ == "__main__":
    main()
