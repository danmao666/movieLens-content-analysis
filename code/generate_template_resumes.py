from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


PROJECT_DIR = Path(__file__).resolve().parents[1]
RESUME_DIR = PROJECT_DIR / "简历"
OUTPUT_DIR = PROJECT_DIR / "days"
PHOTO_FILE = RESUME_DIR / "证件照_简历版.jpg"

NAME = "张博能"
PHONE = "18974648806"
EMAIL = "1786013324@qq.com"
LOCATION = "湘潭"
EDUCATION = "湖南科技大学｜计算机科学与工程｜本科在读｜预计2028年6月毕业"
ENGLISH = "大学英语六级（CET-6）：504分"

NAVY = "244A69"
BLUE = (36, 74, 105)
GRAY = (85, 85, 85)
LIGHT_BLUE = "D9EAF7"


def model_metrics():
    comparison_file = OUTPUT_DIR / "ml_day18_model_comparison.csv"
    pytorch_file = OUTPUT_DIR / "ml_day17_pytorch_metrics.csv"
    forest_auc, forest_f1, pytorch_auc = "0.860", "0.719", "0.835"
    if comparison_file.exists():
        comparison = pd.read_csv(comparison_file)
        row = comparison.loc[
            comparison["model"] == "random_forest_with_enhanced_features"
        ].iloc[0]
        forest_auc = f"{row['roc_auc']:.3f}"
        forest_f1 = f"{row['f1']:.3f}"
    if pytorch_file.exists():
        row = pd.read_csv(pytorch_file).iloc[0]
        pytorch_auc = f"{row['roc_auc']:.3f}"
    return forest_auc, forest_f1, pytorch_auc


FOREST_AUC, FOREST_F1, PYTORCH_AUC = model_metrics()


CONFIGS = {
    "数据分析": {
        "filename": "张博能_数据分析实习_视觉最终版.docx",
        "target": "求职方向：数据分析实习",
        "side_role": "数据分析 / 数据分析实习",
        "summary": "计算机科学与工程本科在读，具备 SQL、Python、Excel 指标分析和基础机器学习项目经验，主要求职数据分析实习。",
        "project": [
            "面对2500万条评分数据，设计用户—内容—行为—日期四表模型；用 DuckDB 完成20+项SQL分析查询，覆盖留存、漏斗、Top N、窗口函数和异常检测。",
            "用 pandas 完成数据质量检查和 EDA；发现高活跃分群平均行为数约430，低活跃分群约26，高活跃约为低活跃16倍，并提出分群运营建议。",
            f"按时间切分构造未来30天活跃预测任务；Random Forest 测试 ROC-AUC {FOREST_AUC}、F1 {FOREST_F1}，完成数据泄漏审计和误判样本分析。",
            "用 PySpark 完成抽样数据的 groupBy、join、窗口排名和分区分析，输出物理执行计划。",
        ],
        "skills": [
            "熟练：SQL（多表关联、CTE、窗口函数、留存、漏斗、Top N、异常检测）；Python（pandas、NumPy、EDA、可视化）",
            "掌握：scikit-learn（逻辑回归、随机森林、时间切分、模型评估）；Excel（数据透视、指标看板）",
            "了解：PyTorch（MLP训练循环）；PySpark（DataFrame、groupBy、join、窗口函数）；Transformer（预训练模型调用）",
        ],
    },
    "内容增长": {
        "filename": "张博能_内容游戏增长分析方向_大字版.docx",
        "target": "求职方向：内容分析｜游戏增长｜用户运营｜产品运营",
        "side_role": "内容增长 / 用户运营",
        "summary": "计算机科学与工程本科在读，围绕公开内容数据完成用户活跃、内容表现、留存、分层、推荐策略和 A/B 实验分析。",
        "project": [
            "围绕“用户为什么不活跃、哪些内容带来高活跃、指标为何波动”建立业务问题和指标字典，明确评分行为只是活跃代理。",
            "使用 SQL 计算 DAU、用户频次、D1/D7/D30 留存、月度 cohort、内容活跃用户和 Top 10 内容集中度，并用 NTILE 完成用户分层。",
            "制作内容排名、用户分层、留存热力图、DAU趋势和集中度图表；将数据发现转化为内容推荐、用户召回和异常排查建议。",
            "设计推荐策略 A/B 实验，以用户为分流单位、七日评分活跃率为主指标，内容多样性和评分质量为护栏指标。",
        ],
        "skills": [
            "增长分析：DAU、留存、cohort、用户频次、用户分层、内容活跃度",
            "内容分析：内容排名、类型表现、长尾覆盖、内容集中度",
            "实验设计：假设、分流单位、主指标、护栏指标、显著性、置信区间",
            "工具：SQL/DuckDB、Python/pandas、Excel、Matplotlib、Seaborn、scikit-learn",
        ],
    },
    "AI数据应用": {
        "filename": "张博能_AI数据应用方向_大字版.docx",
        "target": "求职方向：AI/数据应用助理｜产品/解决方案助理｜数据开发助理",
        "side_role": "AI / 数据应用助理",
        "summary": "计算机科学与工程本科在读，完成 SQL 数据建模、Python 数据处理、基础机器学习、PyTorch 和 PySpark 项目，能够把业务问题转化为可执行的数据流程。",
        "project": [
            "使用 DuckDB 处理 MovieLens 25M 的 25,000,095 条评分行为、162,541 个用户和 59,047 个活跃内容，建立可复用 SQL 指标表。",
            "使用 pandas/NumPy 完成缺失值、重复值、异常值检查、特征构造、EDA和可视化，并用 Excel 输出结果看板。",
            f"完成 Logistic Regression、PyTorch MLP 和 Random Forest 对比；PyTorch 测试 ROC-AUC {PYTORCH_AUC}，Random Forest 测试 ROC-AUC {FOREST_AUC}。",
            "使用 PySpark 完成抽样数据的 groupBy、join、窗口排名、分区和物理计划分析；使用 MovieLens tags 完成 TF-IDF 文本分类基线。",
        ],
        "skills": [
            "Python：pandas、NumPy、scikit-learn、PyTorch基础、PySpark基础",
            "SQL：DuckDB、CTE、窗口函数、数据质量、指标计算、留存和异常检测",
            "机器学习：时间切分、特征工程、分类、ROC-AUC/F1、误判和泄漏检查",
            "应用能力：需求拆解、指标定义、Excel看板、项目报告、A/B方案",
        ],
    },
}


def set_font(run, size=9, bold=False, color=None):
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def shade_cell(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def remove_table_borders(table):
    properties = table._tbl.tblPr
    borders = properties.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "nil")


def paragraph(cell, text="", size=9.0, bold=False, color=None, after=2, align=None):
    first = cell.paragraphs[0]
    has_drawing = bool(first._p.xpath(".//w:drawing"))
    has_picture = bool(first._p.xpath(".//w:pict"))
    p = cell.add_paragraph() if first.text or has_drawing or has_picture else first
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.0
    if text:
        run = p.add_run(text)
        set_font(run, size=size, bold=bold, color=color)
    return p


def sidebar_heading(cell, text):
    p = paragraph(cell, text, size=11.0, bold=True, color=(255, 255, 255), after=5)
    return p


def add_skill_bar(cell, label, blocks):
    p = cell.add_paragraph()
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing = 1.0
    r1 = p.add_run(label)
    set_font(r1, size=8.6, color=(255, 255, 255))
    bar = cell.add_table(rows=1, cols=1)
    bar.autofit = False
    remove_table_borders(bar)
    bar_cell = bar.rows[0].cells[0]
    bar_width = 0.62 * blocks
    bar.columns[0].width = Cm(bar_width)
    bar_cell.width = Cm(bar_width)
    set_cell_margins(bar_cell, top=35, start=0, bottom=35, end=0)
    shade_cell(bar_cell, "B8D2E3")
    bar_paragraph = bar_cell.paragraphs[0]
    bar_paragraph.paragraph_format.space_after = Pt(0)
    bar_paragraph.paragraph_format.line_spacing = 0.2
    spacer = cell.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)
    spacer.paragraph_format.line_spacing = 0.3


def section_header(cell, icon, title):
    p = cell.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.05
    icon_run = p.add_run(f"{icon}  ")
    set_font(icon_run, size=14, bold=True, color=BLUE)
    title_run = p.add_run(title)
    set_font(title_run, size=14, bold=True, color=BLUE)
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "3")
    bottom.set(qn("w:color"), NAVY)
    border.append(bottom)
    p._p.get_or_add_pPr().append(border)


def right_text(cell, text, size=9.4, bold=False, color=GRAY, after=3):
    p = cell.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.08
    run = p.add_run(text)
    set_font(run, size=size, bold=bold, color=color)
    return p


def right_bullet(cell, text):
    p = cell.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Cm(0.35)
    p.paragraph_format.first_line_indent = Cm(-0.16)
    p.paragraph_format.space_after = Pt(2.5)
    p.paragraph_format.line_spacing = 1.08
    run = p.add_run(text)
    set_font(run, size=9.2, color=GRAY)
    return p


def build_resume(config):
    document = Document()
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(0.55)
    section.bottom_margin = Cm(0.55)
    section.left_margin = Cm(0.45)
    section.right_margin = Cm(0.45)

    outer = document.add_table(rows=1, cols=2)
    outer.autofit = False
    outer.rows[0].height = Cm(28.25)
    outer.rows[0].height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
    outer.columns[0].width = Cm(4.1)
    outer.columns[1].width = Cm(15.95)
    remove_table_borders(outer)
    left, right = outer.rows[0].cells
    left.width = Cm(4.1)
    right.width = Cm(15.95)
    left.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    right.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    shade_cell(left, NAVY)
    set_cell_margins(left, top=180, start=250, bottom=150, end=250)
    set_cell_margins(right, top=120, start=360, bottom=100, end=260)

    photo_p = left.paragraphs[0]
    photo_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if PHOTO_FILE.exists():
        photo_p.add_run().add_picture(
            str(PHOTO_FILE), width=Cm(2.75), height=Cm(3.55)
        )
    paragraph(left, "", after=5)
    sidebar_heading(left, "个人信息")
    for item in (
        "▣  本科在读",
        "⌖  湘潭",
        f"☎  {PHONE}",
        f"✉  {EMAIL}",
    ):
        paragraph(left, item, size=8.8, color=(255, 255, 255), after=5)

    sidebar_heading(left, "技能概览")
    add_skill_bar(left, "SQL / 数据分析", 5)
    add_skill_bar(left, "Python / pandas", 5)
    add_skill_bar(left, "Excel / 指标看板", 4)
    add_skill_bar(left, "机器学习 / Spark基础", 2)

    sidebar_heading(left, "补充信息")
    paragraph(left, "CET-6：504分", size=8.8, color=(255, 255, 255), after=4)
    paragraph(left, config["side_role"], size=8.8, color=(210, 225, 235), after=3)
    paragraph(left, "公开数据项目\n本科在读", size=8.4, color=(210, 225, 235), after=2)

    # Clear the default empty paragraph in the right cell.
    right.paragraphs[0].text = ""
    title_p = right.paragraphs[0]
    title_p.paragraph_format.space_after = Pt(1)
    title_run = title_p.add_run(NAME)
    set_font(title_run, size=27, bold=True, color=BLUE)
    target_p = right.add_paragraph()
    target_p.paragraph_format.space_after = Pt(8)
    target_run = target_p.add_run(config["target"])
    set_font(target_run, size=11.5, color=GRAY)

    section_header(right, "●", "教育背景")
    right_text(right, "2025.09—2028.06（预计）    湖南科技大学    计算机科学与工程    本科在读", size=9.5, bold=True)
    right_text(right, "英语六级 504", size=9.2, after=5)

    section_header(right, "●", "项目经历")
    right_text(right, "泛文娱内容平台增长与推荐分析｜MovieLens 25M公开数据模拟项目", size=10.2, bold=True, color=BLUE, after=4)
    for item in config["project"]:
        right_bullet(right, item)

    section_header(right, "●", "专业能力")
    for item in config["skills"]:
        right_bullet(right, item)

    section_header(right, "●", "其他")
    right_bullet(right, "完成 100+ 道 SQL 题；完成 90 分钟 SQL 限时测试（留存、漏斗、Top N、异常）")
    right_bullet(right, "GitHub：github.com/xxx（完整项目代码和文档）")

    section_header(right, "●", "自我评价")
    right_text(
        right,
        "完成1个端到端公开数据项目，能独立完成SQL建模、Python分析、Excel看板和A/B实验设计。对公开数据中的代理指标和局限保持诚实说明，不把模拟结果包装成真实业务成果。",
        size=9.2,
        after=3,
    )
    return document


def main():
    RESUME_DIR.mkdir(parents=True, exist_ok=True)
    config = CONFIGS["数据分析"]
    path = RESUME_DIR / config["filename"]
    document = build_resume(config)
    document.save(path)
    generated = [path]
    print("Generated template resumes:")
    for path in generated:
        print(path)


if __name__ == "__main__":
    main()
